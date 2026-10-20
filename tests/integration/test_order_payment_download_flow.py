"""
End-to-end integration test following the real user journey through the actual
production handlers (not re-implemented logic): wizard data -> presentation
generation -> order saved to DB -> payment confirmed -> file(s) downloaded.

generate_and_preview_presentation() hardcodes its output directory relative to
its own module __file__ (bot/handlers/create_pres.py -> <project_root>/storage/<uuid>),
with no injectable storage path. To keep this test from writing into the real
project's storage/ folder, we temporarily point the module's __file__ at a path
inside pytest's own tmp_path sandbox, which is the only non-invasive way to
redirect it without touching production code.
"""
import os
import pytest
from pptx import Presentation
from bot.handlers import create_pres, payment as payment_module
from tests.fixtures.telegram_fakes import FakeCallbackQuery, FakeUser, FakeFSMContext
from database import repository


@pytest.fixture
def sandboxed_create_pres(monkeypatch, tmp_path):
    fake_module_file = os.path.join(str(tmp_path), "project_root", "bot", "handlers", "create_pres.py")
    monkeypatch.setattr(create_pres, "__file__", fake_module_file)
    return os.path.join(str(tmp_path), "project_root", "storage")


@pytest.mark.asyncio
async def test_full_user_journey_from_wizard_to_download(
    patched_handlers_db, sandboxed_create_pres, no_openai_key, monkeypatch
):
    user_id = 111
    state = FakeFSMContext()
    await state.update_data(
        template_id="05_nature_organic",
        subject="Biologiya",
        topic="Fotosintez jarayoni",
        slide_count=8,
        author_name="Jasur Aliyev",
        institution="TATU",
        language="uz",
        additional_notes="",
        total_price=8 * 350,
    )
    confirm_cb = FakeCallbackQuery(data="confirm_pres_temp_order", from_user=FakeUser(id=user_id))

    # Step 1: user confirms -> presentation is generated and an order is saved.
    await create_pres.generate_and_preview_presentation(confirm_cb, state)

    assert confirm_cb.message.calls_for("answer_media_group"), "2 preview photos must be sent"
    payment_offer = confirm_cb.message.calls_for("answer")[-1]
    assert "Narxi" in payment_offer["text"]

    async with patched_handlers_db() as session:
        from sqlalchemy import select
        from database.models import PresentationOrder
        result = await session.execute(select(PresentationOrder).where(PresentationOrder.user_id == user_id))
        order = result.scalars().one()

    assert order.status == "PREVIEW_READY"
    assert order.slide_count == 8
    assert order.template_id == "05_nature_organic"
    assert order.pptx_path.startswith(sandboxed_create_pres)  # proves it never touched the real storage/ dir
    # the PDF and PPTX are built in the background right after the previews were sent
    from services.presentation import ensure_files
    await ensure_files(os.path.dirname(order.pptx_path))
    assert os.path.exists(order.pptx_path)
    assert os.path.exists(order.pdf_path)

    prs = Presentation(order.pptx_path)
    assert len(prs.slides) == 8

    # Step 2: payment confirmation (DEMO mode).
    monkeypatch.setattr("config.settings.PAYMENT_ENV", "DEMO")
    check_cb = FakeCallbackQuery(data=f"check_pay_{order.id}", from_user=FakeUser(id=user_id))
    await payment_module.process_check_payment(check_cb)

    async with patched_handlers_db() as session:
        updated_order = await repository.get_presentation_order(session, order.id)
        assert updated_order.status == "PAID"

    # Step 3: user downloads both formats -- must be the SAME files generated in
    # step 1, not freshly regenerated placeholders.
    dl_cb = FakeCallbackQuery(data=f"dl_both_{order.id}", from_user=FakeUser(id=user_id))
    await payment_module.process_download_both(dl_cb)

    docs = dl_cb.message.calls_for("answer_document")
    assert len(docs) == 2
    sent_paths = {d["document"].path for d in docs}
    assert order.pptx_path in sent_paths
    assert order.pdf_path in sent_paths


@pytest.mark.asyncio
async def test_a_different_user_cannot_pay_for_or_download_someone_elses_order(
    patched_handlers_db, sandboxed_create_pres, no_openai_key
):
    owner_id, intruder_id = 111, 999
    state = FakeFSMContext()
    await state.update_data(
        template_id="01_midnight_tech", subject="S", topic="T", slide_count=5,
        author_name="A", institution="I", language="uz", additional_notes="", total_price=1750,
    )
    confirm_cb = FakeCallbackQuery(data="confirm_pres_temp_order", from_user=FakeUser(id=owner_id))
    await create_pres.generate_and_preview_presentation(confirm_cb, state)

    async with patched_handlers_db() as session:
        from sqlalchemy import select
        from database.models import PresentationOrder
        result = await session.execute(select(PresentationOrder).where(PresentationOrder.user_id == owner_id))
        order = result.scalars().one()

    intruder_check_cb = FakeCallbackQuery(data=f"check_pay_{order.id}", from_user=FakeUser(id=intruder_id))
    await payment_module.process_check_payment(intruder_check_cb)
    assert "Access Denied" in intruder_check_cb.calls_for("answer")[0]["text"]

    async with patched_handlers_db() as session:
        still_unpaid = await repository.get_presentation_order(session, order.id)
        assert still_unpaid.status != "PAID"

    intruder_dl_cb = FakeCallbackQuery(data=f"dl_both_{order.id}", from_user=FakeUser(id=intruder_id))
    await payment_module.process_download_both(intruder_dl_cb)
    assert intruder_dl_cb.message.calls_for("answer_document") == []


@pytest.mark.asyncio
async def test_production_mode_offers_click_and_payme_instead_of_mock(
    patched_handlers_db, sandboxed_create_pres, no_openai_key, monkeypatch
):
    monkeypatch.setattr("config.settings.PAYMENT_ENV", "PRODUCTION")
    state = FakeFSMContext()
    await state.update_data(
        template_id="01_midnight_tech", subject="S", topic="T", slide_count=5,
        author_name="A", institution="I", language="uz", additional_notes="", total_price=1750,
    )
    confirm_cb = FakeCallbackQuery(data="confirm_pres_temp_order", from_user=FakeUser(id=111))

    await create_pres.generate_and_preview_presentation(confirm_cb, state)

    payment_offer = confirm_cb.message.calls_for("answer")[-1]
    callback_datas = [btn.callback_data for row in payment_offer["reply_markup"].inline_keyboard for btn in row if btn.callback_data]
    assert any(cd.startswith("pay_click_") for cd in callback_datas)
    assert any(cd.startswith("pay_payme_") for cd in callback_datas)
    assert not any(cd.startswith("pay_checkout_") for cd in callback_datas)  # no mock button in production


@pytest.mark.asyncio
async def test_full_production_journey_via_click_webhook(
    patched_handlers_db, sandboxed_create_pres, no_openai_key, monkeypatch
):
    """The complete real-money-shaped path: wizard -> Click checkout link shown ->
    Click's Prepare/Complete webhook calls arrive -> order flips to PAID -> user
    downloads via the bot, all through the actual production code paths."""
    from services.payment import click_webhook
    monkeypatch.setattr("config.settings.PAYMENT_ENV", "PRODUCTION")
    monkeypatch.setattr("config.settings.CLICK_SERVICE_ID", "20000")
    monkeypatch.setattr("config.settings.CLICK_MERCHANT_ID", "12345")
    monkeypatch.setattr("config.settings.CLICK_SECRET_KEY", "integration-test-secret")

    state = FakeFSMContext()
    await state.update_data(
        template_id="01_midnight_tech", subject="S", topic="T", slide_count=5,
        author_name="A", institution="I", language="uz", additional_notes="", total_price=1750,
    )
    confirm_cb = FakeCallbackQuery(data="confirm_pres_temp_order", from_user=FakeUser(id=111))
    await create_pres.generate_and_preview_presentation(confirm_cb, state)

    async with patched_handlers_db() as session:
        from sqlalchemy import select
        from database.models import PresentationOrder
        result = await session.execute(select(PresentationOrder).where(PresentationOrder.user_id == 111))
        order = result.scalars().one()

    click_cb = FakeCallbackQuery(data=f"pay_click_{order.id}", from_user=FakeUser(id=111))
    await payment_module.process_pay_checkout_click(click_cb)

    # Click now calls our webhook twice (Prepare, then Complete) exactly as it
    # would in production, using the real signature algorithm.
    import hashlib

    def sign(params, action):
        secret = "integration-test-secret"
        if action == 0:
            raw = (str(params["click_trans_id"]) + str(params["service_id"]) + secret
                   + str(params["merchant_trans_id"]) + str(params["amount"]) + str(action) + str(params["sign_time"]))
        else:
            raw = (str(params["click_trans_id"]) + str(params["service_id"]) + secret
                   + str(params["merchant_trans_id"]) + str(params["merchant_prepare_id"])
                   + str(params["amount"]) + str(action) + str(params["sign_time"]))
        return hashlib.md5(raw.encode()).hexdigest()

    prepare_params = {
        "click_trans_id": "7001", "service_id": "20000", "merchant_trans_id": order.id,
        "amount": "1750", "action": "0", "sign_time": "2026-01-01 00:00:00",
    }
    prepare_params["sign_string"] = sign(prepare_params, 0)
    prepare_resp = await click_webhook.handle_prepare(prepare_params)
    assert prepare_resp["error"] == 0

    complete_params = {
        "click_trans_id": "7001", "service_id": "20000", "merchant_trans_id": order.id,
        "merchant_prepare_id": prepare_resp["merchant_prepare_id"], "amount": "1750", "action": "1",
        "sign_time": "2026-01-01 00:00:01", "error": "0", "error_note": "",
    }
    complete_params["sign_string"] = sign(complete_params, 1)
    complete_resp = await click_webhook.handle_complete(complete_params)
    assert complete_resp["error"] == 0

    # User comes back and taps "Tasdiqlash" -- status is already PAID from the webhook.
    check_cb = FakeCallbackQuery(data=f"check_pay_{order.id}", from_user=FakeUser(id=111))
    await payment_module.process_check_payment(check_cb)

    dl_cb = FakeCallbackQuery(data=f"dl_both_{order.id}", from_user=FakeUser(id=111))
    await payment_module.process_download_both(dl_cb)
    assert len(dl_cb.message.calls_for("answer_document")) == 2
