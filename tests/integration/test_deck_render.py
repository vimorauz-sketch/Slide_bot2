"""Renders real decks through Chromium: PNG, PDF and editable PPTX from the same HTML."""
import json
import os
import pytest
from pptx import Presentation
from services import ai_service
from services.deck import render_deck
from services.deck.themes import list_family_ids, list_theme_ids

FIXTURE = os.path.join(os.path.dirname(os.path.dirname(__file__)), "fixtures", "fotosintez_content.json")


def _content():
    with open(FIXTURE, encoding="utf-8") as f:
        return json.load(f)


@pytest.mark.parametrize("theme", ["05_nature_organic", "01_midnight_tech", "09_uzbek_heritage"])
async def test_full_sample_deck_renders_without_overflow_in_every_format(theme, storage_dir):
    content = _content()
    res = await render_deck(content, theme, storage_dir)
    n = 1 + len(content["slides"])
    assert len(res.slide_pngs) == n
    assert not [f for f in res.fit if f["overflow"]], res.fit
    from pypdf import PdfReader
    assert len(PdfReader(res.pdf_path).pages) == n
    prs = Presentation(res.pptx_path)
    assert len(prs.slides) == n
    texts = [sh.text_frame.text for sh in prs.slides[2].shapes if sh.has_text_frame]
    assert any("Fotosintez nima?" in t for t in texts)  # real, editable text, not just a picture
    notes = prs.slides[2].notes_slide.notes_text_frame.text
    assert "taʼkidlang" in notes or "ta’kidlang" in notes


@pytest.mark.parametrize("theme", list_family_ids() + list_theme_ids()[20::9] + ["150_apollo~emerald", "12_pastel_playful~pomegranate"])
async def test_templates_render_a_mock_deck(theme, storage_dir, no_openai_key):
    content = ai_service.generate_mock_presentation({"topic": "Suv aylanishi", "subject": "Geografiya", "slide_count": 8,
                                                     "author_name": "A", "institution": "B"})
    res = await render_deck(content, theme, storage_dir, want_pdf=False, want_pptx=False)
    assert len(res.slide_pngs) == 8
    assert not [f for f in res.fit if f["overflow"]]


async def test_custom_design_spec_renders(storage_dir):
    spec = {"background_color": "#101820", "primary_color": "#FEE715", "accent_color": "#00A8E8", "text_color": "#FFFFFF"}
    content = ai_service.generate_mock_presentation({"topic": "T", "slide_count": 5})
    res = await render_deck(content, spec, storage_dir, want_pdf=False, want_pptx=False)
    assert len(res.slide_pngs) == 5


async def test_uzbek_text_is_normalized_in_the_output(storage_dir):
    content = {"language": "uz", "title_slide": {"title": "O'zbekiston"},
               "slides": [{"layout": "text_image", "title": "Ta'lim", "lead": "o'quv"}]}
    res = await render_deck(content, "02", storage_dir, want_pdf=False, want_pptx=False, want_pngs=False)
    html = open(res.html_path, encoding="utf-8").read()
    assert "O‘zbekiston" in html and "Ta’lim" in html and "o‘quv" in html


async def test_long_text_is_shrunk_or_trimmed_instead_of_overflowing(storage_dir):
    long = "juda uzun matn " * 30
    content = {"language": "uz", "title_slide": {"title": "T"}, "slides": [
        {"layout": "cards", "title": long, "lead": long,
         "items": [{"icon": "leaf", "title": long, "text": long} for _ in range(4)]}]}
    res = await render_deck(content, "06", storage_dir, want_pdf=False, want_pptx=False)
    assert not [f for f in res.fit if f["overflow"]]


async def test_photo_template_puts_a_photo_on_every_card(storage_dir, monkeypatch, tmp_path):
    from PIL import Image
    folder = tmp_path / "photos"
    folder.mkdir()
    for k in range(8):
        Image.new("RGB", (1600, 1200), (30 * k, 140, 90)).save(folder / f"p{k}.jpg")
    monkeypatch.setattr("config.settings.LOCAL_PHOTO_DIR", str(folder))
    content = {"language": "uz", "title_slide": {"title": "Biosfera modeli", "subject": "Biologiya"},
               "slides": [{"layout": "cards", "title": "Biosfera qatlamlari",
                           "items": [{"title": t, "text": "Matn", "photo": q} for t, q in
                                     (("Atmosfera", "sky clouds"), ("Gidrosfera", "ocean"), ("Litosfera", "mountains"))]},
                          {"layout": "summary", "title": "Xulosa", "points": ["A", "B", "C"]}]}
    res = await render_deck(content, "201", storage_dir, want_pdf=False, want_pptx=True)
    html = open(res.html_path, encoding="utf-8").read()
    assert html.count('class="card-photo"') == 3          # real photos, not the art fallback
    assert "photo-mode" in html and "layout-summary with-bleed" in html
    assert not [f for f in res.fit if f["overflow"]]
