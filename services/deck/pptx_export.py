"""
Hybrid PPTX export.

Each slide = one background picture (the rendered slide with its text hidden: shapes,
gradients, photos, charts, icons) + native, editable PowerPoint text boxes placed exactly
where the browser drew each text, with the browser's own line breaks. The result looks
like the PDF and the text can still be edited in PowerPoint, WPS or Google Slides.
"""
from __future__ import annotations

import re
from typing import Any, Dict, List

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 6350  # EMU per CSS px on a 1920 px wide, 13.333 in slide
SLIDE_W, SLIDE_H = 1920 * PX, 1080 * PX


def _rgb(css: str, fallback: str = "000000") -> RGBColor:
    m = re.match(r"rgba?\((\d+),\s*(\d+),\s*(\d+)", css or "")
    if m:
        return RGBColor(*(int(x) for x in m.groups()))
    h = (css or "").lstrip("#")
    if len(h) == 6:
        return RGBColor.from_string(h.upper())
    return RGBColor.from_string(fallback)


def _alpha(css: str) -> float:
    m = re.match(r"rgba\(\d+,\s*\d+,\s*\d+,\s*([\d.]+)\)", css or "")
    return float(m.group(1)) if m else 1.0


def _add_text(slide, item: Dict[str, Any], theme: Dict[str, Any]) -> None:
    x, y, w, h = item["x"], item["y"], item["w"], item["h"]
    align = item.get("align", "left")
    # A little extra width absorbs rounding differences between renderers; the browser's
    # own line breaks are kept as separate paragraphs, so nothing re-wraps.
    extra = max(12.0, w * 0.06)
    if align == "center":
        x -= extra / 2
    elif align in ("right", "end"):
        x -= extra
    w += extra
    x = max(0.0, x)
    w = min(w, 1920 - x)
    box = slide.shapes.add_textbox(Emu(int(x * PX)), Emu(int(y * PX)), Emu(int(w * PX)), Emu(int((h + 4) * PX)))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP

    color = item.get("color", "")
    if item.get("gradientText") or _alpha(color) < 0.05:
        color = theme["palette"].get("primary", "#000000")
    a = _alpha(color)
    scale = min(1.0, max(0.5, float(item.get("scale", 1.0))))
    size_pt = max(6.0, round(item["fontSize"] * 0.5 * scale * 2) / 2)
    pa = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)

    for i, line in enumerate(item["lines"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = pa
        p.line_spacing = Pt(item["lineHeight"] * 0.5)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        f = run.font
        f.size = Pt(size_pt)
        f.bold = item.get("fontWeight", 400) >= 600
        f.italic = bool(item.get("italic"))
        f.name = item.get("font") or "Arial"
        f.color.rgb = _rgb(color)
        rpr = run._r.get_or_add_rPr()
        if item.get("letterSpacing"):
            rpr.set("spc", str(int(item["letterSpacing"] * 0.5 * scale * 100)))
        if a < 0.99:
            srgb = rpr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(a * 100000))}))


def build_pptx(path: str, backgrounds: List[str], measured: List[List[Dict[str, Any]]],
               slides: List[Dict[str, Any]], theme: Dict[str, Any]) -> str:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    for i, bg in enumerate(backgrounds):
        sl = prs.slides.add_slide(blank)
        pic = sl.shapes.add_picture(bg, 0, 0, SLIDE_W, SLIDE_H)
        pic.name = "Background"
        for item in measured[i] if i < len(measured) else []:
            try:
                _add_text(sl, item, theme)
            except Exception:  # one odd text box should never break the whole file
                continue
        notes = (slides[i].get("notes") if i < len(slides) else "") or ""
        if notes:
            sl.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path
