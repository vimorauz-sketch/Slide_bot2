"""
Editable Word (.docx) version of a referat.

The title page's decoration (frame, ornament, colour bands, cover photo) is a full-page picture
behind the text; every word on the page stays editable. The text sits in a borderless table with
exact row heights, so it lands where the PDF has it. Inner pages follow the academic standard.
"""
from __future__ import annotations

import copy
from typing import Any, Dict, Optional

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from services.deck import fonts

# Word names of library fonts that have a standard Windows twin with the same metrics
WINDOWS_TWINS = {"g:tinos": "Times New Roman", "g:arimo": "Arial"}

ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}


def word_font(font_id: str) -> str:
    return WINDOWS_TWINS.get(font_id) or fonts.family_name(font_id) or "Times New Roman"


def _rgb(hex_: str) -> RGBColor:
    h = hex_.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(run, name: str, size: Optional[float] = None, bold: Optional[bool] = None, color: Optional[str] = None,
              italic: Optional[bool] = None, spacing: Optional[float] = None) -> None:
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = _rgb(color)
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        rpr.append(sp)


def _para(container, text: str = "", *, font: str, size: float, bold=False, color=None, align=None, italic=False,
          spacing=None, before: float = 0, after: float = 0, line: float = 1.15, indent_first: Optional[float] = None,
          left: Optional[float] = None):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if indent_first is not None:
        pf.first_line_indent = Cm(indent_first)
    if left is not None:
        pf.left_indent = Mm(left)
    if text:
        _set_font(p.add_run(text), font, size, bold, color, italic, spacing)
    return p


def _no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _cell_margins_zero(table) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _shade(cell, hex_: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_.lstrip("#"))
    tc_pr.append(shd)


def _behind_text(run, width_emu: int, height_emu: int) -> None:
    """Turns the run's inline picture into a full-page picture behind the text, anchored to the page."""
    inline = run._r.find(".//" + qn("wp:inline"))
    if inline is None:
        return
    graphic = inline.find(qn("a:graphic"))
    doc_pr = inline.find(qn("wp:docPr"))
    anchor = OxmlElement("wp:anchor")
    for k, v in (("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"), ("simplePos", "0"),
                 ("relativeHeight", "0"), ("behindDoc", "1"), ("locked", "1"), ("layoutInCell", "0"),
                 ("allowOverlap", "1")):
        anchor.set(k, v)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for tag, rel in (("wp:positionH", "page"), ("wp:positionV", "page")):
        pos = OxmlElement(tag)
        pos.set("relativeFrom", rel)
        off = OxmlElement("wp:posOffset")
        off.text = "0"
        pos.append(off)
        anchor.append(pos)
    extent = OxmlElement("wp:extent")
    extent.set("cx", str(width_emu))
    extent.set("cy", str(height_emu))
    anchor.append(extent)
    eff = OxmlElement("wp:effectExtent")
    for k in ("l", "t", "r", "b"):
        eff.set(k, "0")
    anchor.append(eff)
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(copy.deepcopy(doc_pr))
    anchor.append(OxmlElement("wp:cNvGraphicFramePr"))
    anchor.append(copy.deepcopy(graphic))
    inline.getparent().replace(inline, anchor)


def _page_number_footer(section, font: str, color: str) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _set_font(run, font, 11, color=color)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)


def build_docx(path: str, vm: Dict[str, Any], template: Dict[str, Any], cover_background: Optional[str],
               images: Dict[str, str], running: str = "") -> str:
    """vm: services.referat.render.view_model(); images: section key -> local file path."""
    p, t = template["palette"], template
    tf, bf = word_font(t["fonts"]["title"]), word_font(t["fonts"]["body"])
    body_cat = (fonts.library_font(t["fonts"]["body"]) or {}).get("category")
    body_size = 13 if body_cat in ("modern", "tech") else 14
    m, L = vm["m"], vm["L"]
    g = dict(t["layout"])  # title page geometry in mm, shared with referat.html.j2
    tone = g.get("tone", "dark")
    light = tone == "light"

    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = bf
    normal.font.size = Pt(body_size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), bf)
    normal.paragraph_format.space_after = Pt(0)
    doc.core_properties.title = m.get("topic", "")
    doc.core_properties.author = m.get("author", "")
    doc.core_properties.subject = m.get("subject", "")

    # ------------------------------------------------------------ title page
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(g["top"]), Mm(12)
    sec.left_margin, sec.right_margin = Mm(g["left"]), Mm(g["right"])
    sec.header_distance = sec.footer_distance = Mm(5)

    content_w = 210 - g["left"] - g["right"]
    # the last rows leave ~12 mm free: Word needs room for the section break paragraph after the table
    city_top = max(g["info"] + 22, 297 - g["city"] - 9)
    rows_h = [g["mid"] - g["top"], g["info"] - g["mid"], city_top - g["info"], 14]
    table = doc.add_table(rows=4, cols=1)
    _no_borders(table)
    _cell_margins_zero(table)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row, h in zip(table.rows, rows_h):
        row.height = Mm(h)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        row.cells[0].width = Mm(content_w)
    align = ALIGN[g["align"]]
    top_color = "#FFFFFF" if tone in ("light", "top_light") else p["ink"]
    ink = "#FFFFFF" if light else p["ink"]
    accent = "#FFFFFF" if light else p["accent"]
    kind_color = p["accent2"] if light else (p["ink"] if t["cover"] == "magazine" else p["accent"])
    c0, c1, c2, c3 = (r.cells[0] for r in table.rows)
    c0.paragraphs[0].text = ""
    if m.get("ministry"):
        _para(c0, m["ministry"], font=bf, size=11, bold=True, color=top_color, align=align, line=1.2)
    if m.get("institution"):
        _para(c0, m["institution"], font=bf, size=12.5, bold=True, color=top_color, align=align, before=8, line=1.2)
    if m.get("faculty"):
        _para(c0, m["faculty"], font=bf, size=11.5, color=top_color, align=align, before=3)
    _delete_first_empty(c0)

    kind_size = {"minimal": 15, "magazine": 60, "card": 30}.get(t["cover"], 36)
    if t["cover"] != "split":
        _para(c1, m.get("kind", ""), font=tf, size=kind_size, bold=True, color=kind_color, align=align,
              spacing=kind_size * (0.3 if t["cover"] == "minimal" else 0.14), line=1.0)
    if m.get("subject"):
        sp = _para(c1, "", font=bf, size=13, align=align, before=16)
        _set_font(sp.add_run(f"{L['subject']}: "), bf, 13, color=ink)
        _set_font(sp.add_run(m["subject"]), bf, 13, bold=True, color=accent)
    _para(c1, f"{L['topic']}:".upper(), font=bf, size=12, bold=True, color=accent, align=align, before=14, spacing=1)
    script = t["fonts"]["title"] in ("g:greatvibes", "g:dancingscript", "g:cookie", "g:parisienne")
    _para(c1, f"«{m.get('topic', '')}»", font=bf if script else tf, size={"minimal": 26, "magazine": 24}.get(t["cover"], 19), bold=True,
          color=ink, align=align, before=4, line=1.2)
    _delete_first_empty(c1)
    if t["cover"] == "split":
        _para(c1, m.get("kind", ""), font=tf, size=30, bold=True, color=p["accent"], align=align, before=18, spacing=6)

    info_indent = max(0, g["info_left"] - g["left"])
    for key, label in (("author", L["author"]), ("group", m.get("group_label") or L["group"]), ("teacher", L["teacher"])):
        if m.get(key):
            ip = _para(c2, "", font=bf, size=12.5, left=info_indent, line=1.5)
            ip.paragraph_format.tab_stops.add_tab_stop(Mm(info_indent + 28))
            _set_font(ip.add_run(f"{label}:\t"), bf, 12.5, bold=True, color=accent)
            _set_font(ip.add_run(m[key]), bf, 12.5, color=ink)
    _delete_first_empty(c2)
    city = " – ".join(x for x in (m.get("city"), m.get("year")) if x)
    city_color = "#FFFFFF" if (light or g.get("city_tone") == "light") else p["ink"]
    cp = _para(c3, city, font=bf, size=12.5, bold=True, color=city_color, align=align)
    _delete_first_empty(c3)
    if cover_background:  # anchored to the page, so it can live in any paragraph of the title page
        run = cp.add_run()
        run.add_picture(cover_background, width=Mm(210), height=Mm(297))
        _behind_text(run, Mm(210), Mm(297))

    # ------------------------------------------------------------ inner pages
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.top_margin, body.bottom_margin = Mm(20), Mm(20)
    body.left_margin, body.right_margin = Mm(30), Mm(15)
    body.header_distance = body.footer_distance = Mm(10)
    _page_number_footer(body, bf, p["ink"])
    _page_decor(body, t.get("page", "none"), p)
    header = body.header
    header.is_linked_to_previous = False
    if running:
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_font(hp.add_run(running.upper()), tf, 8.5, color=p["accent"], spacing=0.5)
    first_header = sec.header
    first_header.is_linked_to_previous = False

    heading_style = t["heading"]
    h_align = WD_ALIGN_PARAGRAPH.CENTER if heading_style in ("caps", "rule") else WD_ALIGN_PARAGRAPH.LEFT
    h_color = {"caps": p["ink"], "marker": p["ink"], "pill": "#FFFFFF", "block": "#FFFFFF"}.get(heading_style, p["accent"])

    def heading(text: str, level: int, page_break: bool):
        hp = doc.add_paragraph()
        if page_break and doc.paragraphs and len(doc.paragraphs) > 1:
            hp.paragraph_format.page_break_before = True
        hp.alignment = h_align if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_before = Pt(0 if level == 1 else 10)
        hp.paragraph_format.space_after = Pt(12 if level == 1 else 6)
        hp.paragraph_format.keep_with_next = True
        hp.style = doc.styles["Heading 1" if level == 1 else "Heading 2"]
        txt = text.upper() if (level == 1 and heading_style in ("caps", "overline")) else text
        _set_font(hp.add_run(txt), tf, 16 if level == 1 else 14, bold=True,
                  color=h_color if level == 1 else (p["ink"] if heading_style in ("caps", "marker") else p["accent"]))
        if level == 1 and heading_style in ("box", "marker"):
            _para_shade(hp, p["soft"])
        if level == 1 and heading_style in ("pill", "block"):
            _para_shade(hp, p["accent"])
        if level == 1 and heading_style in ("bar", "underline", "overline", "numbig"):
            side = {"bar": "left", "underline": "bottom", "overline": "top", "numbig": "bottom"}[heading_style]
            _para_border(hp, side, p["accent2"] if heading_style in ("underline", "numbig") else p["accent"])
        return hp

    def body_par(text: str):
        _para(doc, text, font=bf, size=body_size, color=p["ink"], align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5,
              indent_first=1.25)

    # plan
    pt = _para(doc, L["plan"].upper(), font=tf, size=16, bold=True, color=p["accent"], align=WD_ALIGN_PARAGRAPH.CENTER,
               after=16, spacing=2)
    pt.paragraph_format.keep_with_next = True
    for s in vm["sections"]:
        lp = _para(doc, "", font=bf, size=body_size, line=1.5)
        if s["number"]:
            _set_font(lp.add_run(s["number"] + " "), bf, body_size, bold=True, color=p["accent"])
        _set_font(lp.add_run(s["title"]), bf, body_size, color=p["ink"])
        for sub in s["subsections"]:
            sp = _para(doc, "", font=bf, size=body_size, line=1.5, left=10)
            if sub["number"]:
                _set_font(sp.add_run(sub["number"] + " "), bf, body_size, bold=True, color=p["accent"])
            _set_font(sp.add_run(sub["title"]), bf, body_size, color=p["ink"])
    _para(doc, L["references"], font=bf, size=body_size, color=p["ink"], line=1.5)

    def figure(key: str, img: Dict[str, Any]):
        path = images.get(key)
        if not path:
            return
        fp = doc.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(10)
        fp.paragraph_format.keep_with_next = True
        try:
            fp.add_run().add_picture(path, width=Cm(12))
        except Exception:  # noqa: BLE001 - a broken picture must not break the document
            return
        cp = _para(doc, "", font=bf, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
        _set_font(cp.add_run(f"{img['n']}-{L['figure']}. "), bf, 11.5, bold=True, color=p["accent"])
        _set_font(cp.add_run(img["caption"]), bf, 11.5, italic=True, color=p["ink"])

    def table_block(tb: Dict[str, Any]):
        cap = _para(doc, "", font=bf, size=11.5, align=WD_ALIGN_PARAGRAPH.RIGHT, before=10, after=4)
        cap.paragraph_format.keep_with_next = True
        _set_font(cap.add_run(f"{tb['n']}-{L['table']}. "), bf, 11.5, bold=True, color=p["accent"])
        _set_font(cap.add_run(tb["caption"]), bf, 11.5, color=p["ink"])
        cols = max(len(tb["head"]), max((len(r) for r in tb["rows"]), default=0))
        tt = doc.add_table(rows=1 + len(tb["rows"]), cols=cols)
        tt.style = doc.styles["Table Grid"]
        for j in range(cols):
            c = tt.rows[0].cells[j]
            c.text = ""
            _set_font(c.paragraphs[0].add_run(tb["head"][j] if j < len(tb["head"]) else ""), bf, 12, bold=True, color=p["accent"])
            _shade(c, p["soft"])
        for i, r in enumerate(tb["rows"], start=1):
            for j in range(cols):
                c = tt.rows[i].cells[j]
                c.text = ""
                _set_font(c.paragraphs[0].add_run(r[j] if j < len(r) else ""), bf, 12, color=p["ink"])
        _para(doc, "", font=bf, size=6)

    for s in vm["sections"]:
        heading((f"{s['number']} " if s["number"] else "") + s["title"], 1, True)
        for para in s["paragraphs"]:
            body_par(para)
        for k, sub in enumerate(s["subsections"]):
            heading((f"{sub['number']} " if sub["number"] else "") + sub["title"], 2, False)
            for para in sub["paragraphs"]:
                body_par(para)
            if sub.get("image"):
                figure(sub["image"]["key"], sub["image"])
            if sub.get("table"):
                table_block(sub["table"])
        if s.get("image"):
            figure(s["image"]["key"], s["image"])
    heading(L["references"], 1, True)
    for i, r in enumerate(vm["references"], start=1):
        rp = _para(doc, "", font=bf, size=body_size, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5, after=4)
        _set_font(rp.add_run(f"{i}. "), bf, body_size, bold=True, color=p["accent"])
        _set_font(rp.add_run(r), bf, body_size, color=p["ink"])
    doc.save(path)
    return path


def _delete_first_empty(cell) -> None:
    ps = cell.paragraphs
    if len(ps) > 1 and not ps[0].text.strip():
        ps[0]._p.getparent().remove(ps[0]._p)


def _para_shade(paragraph, hex_: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_.lstrip("#"))
    ppr.append(shd)


def _page_decor(section, kind: str, p: Dict[str, str]) -> None:
    """Inner page decoration of the template, as Word page borders (measured from the page edge)."""
    sides = {"band_top": {"top": (96, p["accent"])}, "band_bottom": {"bottom": (96, p["accent"])},
             "side_band": {"left": (96, p["accent"])}, "side_line": {"left": (6, p["accent2"])},
             "frame": {s: (6, p["accent"]) for s in ("top", "left", "bottom", "right")},
             "corners": {s: (4, p["accent2"]) for s in ("top", "left", "bottom", "right")}}.get(kind)
    if not sides:
        return
    sect = section._sectPr
    bdr = OxmlElement("w:pgBorders")
    bdr.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            sz, color = sides[side]
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "24" if sz < 50 else "0")
            el.set(qn("w:color"), color.lstrip("#"))
            bdr.append(el)
    # pgBorders goes after w:pgMar / w:pgSz and before w:cols etc.
    anchor = sect.find(qn("w:pgMar"))
    if anchor is not None:
        anchor.addnext(bdr)
    else:
        sect.append(bdr)


def _para_border(paragraph, side: str, hex_: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "24" if side == "left" else "12")
    el.set(qn("w:space"), "8" if side == "left" else "4")
    el.set(qn("w:color"), hex_.lstrip("#"))
    bdr.append(el)
    ppr.append(bdr)
